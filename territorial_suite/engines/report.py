"""Territorial dossier: HTML, PDF and tables.

The report is the written half of the plugin's promise: every statement carries its source,
its date and the nature of the data it rests on. Sections with no data say so explicitly -
"no data" and "no constraint" are different answers, and confusing them is how a knowledge
document turns into a false reassurance.
"""

from __future__ import annotations

import csv
import html
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from ..core import log, measure, settings
from ..core.constants import PLUGIN_NAME, PLUGIN_VERSION
from ..core.errors import EngineError
from ..core.models import AnalysisReport, CadastralRow, SourceResult
from ..core.taxonomy import Taxonomy

CSS = """
body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 10pt; color: #222; }
h1 { font-size: 18pt; margin-bottom: 2pt; }
h2 { font-size: 13pt; margin-top: 16pt; margin-bottom: 4pt;
     border-bottom: 1px solid #999; padding-bottom: 2pt; }
h3 { font-size: 11pt; margin-top: 10pt; margin-bottom: 3pt; }
p, li, td, th { font-size: 9pt; }
.subtitle { color: #555; font-size: 10pt; margin-top: 0; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0; }
th { background: #eee; text-align: left; padding: 3pt; border: 1px solid #bbb; }
td { padding: 3pt; border: 1px solid #ccc; vertical-align: top; }
.small { font-size: 8pt; color: #555; }
.note { background: #f6f6f6; border-left: 3px solid #999; padding: 6pt; margin: 8pt 0; }
.alert-CHECK_REQUIRED { border-left: 3px solid #444; }
.alert-ATTENTION { border-left: 3px solid #777; }
.alert-CARTOGRAPHIC_ISSUE { border-left: 3px solid #aaa; }
.alert-INFO { border-left: 3px solid #ccc; }
.alert { padding: 4pt 8pt; margin: 4pt 0; background: #fafafa; }
.level { font-weight: bold; font-size: 8pt; letter-spacing: 0.5pt; }
.evidence { font-style: italic; color: #555; font-size: 8pt; }
.missing { color: #666; font-style: italic; }
"""

DISCLAIMER = (
    "Il presente documento e' un quadro conoscitivo costruito automaticamente a partire "
    "dai dati cartografici delle fonti indicate. Le informazioni riportate hanno natura "
    "ricognitiva: non costituiscono accertamento dei vincoli ne' certificazione "
    "urbanistica, catastale o ambientale. La verifica dei vincoli e delle prescrizioni "
    "applicabili va effettuata presso gli enti competenti sulla base degli atti vigenti."
)


@dataclass
class ReportOptions:
    """What the dossier should contain."""

    title: str = "Quadro conoscitivo territoriale"
    author: str = ""
    organisation: str = ""
    include_cadastre: bool = True
    include_terrain: bool = True
    include_sources: bool = True
    include_alerts: bool = True
    include_cultural_heritage: bool = True
    max_features_per_source: int = 12
    max_cadastral_rows: int = 400
    max_heritage_rows: int = 40
    language: str = "it"

    @classmethod
    def from_settings(cls) -> "ReportOptions":
        """Build the options from the user settings."""
        return cls(
            author=settings.get("report.author", ""),
            organisation=settings.get("report.organisation", ""),
            include_cadastre=bool(settings.get("report.include_cadastre", True)),
            include_terrain=bool(settings.get("report.include_terrain", True)),
            include_sources=bool(settings.get("report.include_sources", True)),
            include_cultural_heritage=bool(
                settings.get("report.include_cultural_heritage", True)),
            max_heritage_rows=int(settings.get("report.max_heritage_rows", 40)),
            language=settings.get("report.language", "it"),
        )


class ReportEngine:
    """Renders an :class:`AnalysisReport` as HTML/PDF and exports its tables."""

    def __init__(self, report: AnalysisReport,
                 options: Optional[ReportOptions] = None) -> None:
        self.report = report
        self.options = options or ReportOptions.from_settings()
        self.taxonomy = Taxonomy.instance()

    # ------------------------------------------------------------------ HTML

    def to_html(self) -> str:
        """Render the whole dossier as a standalone HTML document."""
        parts: List[str] = [
            f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body>",
            self._header(),
            self._section_area(),
            self._section_admin(),
            self._section_alerts(),
            self._section_cadastre(),
        ]
        parts.extend(self._knowledge_sections())
        parts.append(self._section_cultural_heritage())
        parts.append(self._section_terrain())
        parts.append(self._section_sources())
        parts.append(self._section_metadata())
        parts.append(f"<div class='note'>{html.escape(DISCLAIMER)}</div>")
        parts.append("</body></html>")
        return "\n".join(part for part in parts if part)

    def write_html(self, path: Path) -> Path:
        """Write the HTML dossier."""
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(self.to_html(), encoding="utf-8")
        return path

    def write_pdf(self, path: Path) -> Path:
        """Render the dossier to PDF using Qt's text engine (no extra dependency)."""
        from qgis.PyQt.QtCore import QMarginsF, QSizeF
        from qgis.PyQt.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        writer = QPdfWriter(str(path))
        try:
            writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
            writer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Unit.Millimeter)
            writer.setResolution(300)
        except Exception as exc:  # pragma: no cover - API differences
            log.debug(f"write_pdf: operazione non riuscita ({type(exc).__name__}: {exc})")
        document = QTextDocument()
        document.setHtml(self.to_html())
        try:
            document.setPageSize(QSizeF(writer.width(), writer.height()))
        except Exception as exc:  # pragma: no cover - API differences
            log.debug(f"write_pdf: operazione non riuscita ({type(exc).__name__}: {exc})")
        printer = getattr(document, "print", None) or getattr(document, "print_")
        printer(writer)
        if not path.exists():
            raise EngineError("Il PDF del report non e' stato prodotto")
        return path

    # ------------------------------------------------------------------ sections

    def _header(self) -> str:
        area = self.report.area
        subtitle = area.get("name", "")
        meta = []
        if self.options.organisation:
            meta.append(html.escape(self.options.organisation))
        if self.options.author:
            meta.append(html.escape(self.options.author))
        meta.append(self._format_date(self.report.finished_at or self.report.started_at))
        return (f"<h1>{html.escape(self.options.title)}</h1>"
                f"<p class='subtitle'>{html.escape(subtitle)}</p>"
                f"<p class='small'>{' &middot; '.join(meta)}</p>")

    def _section_area(self) -> str:
        area = self.report.area
        bbox = self._metric_bbox(area)
        rows = [
            ("Denominazione", area.get("name", "")),
            ("Identificativo", area.get("id", "")),
            ("Superficie", measure.format_area(float(area.get("area_m2", 0.0) or 0.0))),
            ("Perimetro", measure.format_length(float(area.get("perimeter_m", 0.0) or 0.0))),
            ("Sistema di riferimento", area.get("crs", "")),
            ("CRS di lavoro", area.get("work_crs", "")),
            ("Ellissoide di calcolo", area.get("ellipsoid", "")),
            (f"Bounding box ({area.get('work_crs', '')})",
             f"X {bbox['min_x']:,.1f} - {bbox['max_x']:,.1f} | "
             f"Y {bbox['min_y']:,.1f} - {bbox['max_y']:,.1f} "
             f"({bbox['width']:,.0f} x {bbox['height']:,.0f} m)"),
            ("Origine", area.get("source_kind", "")),
            ("Data di creazione", self._format_date(area.get("created_at", ""))),
        ]
        return "<h2>1. Area di progetto</h2>" + self._table_two_columns(rows)

    def _section_admin(self) -> str:
        admin = self.report.admin
        if not admin.municipalities:
            return ("<h2>2. Localizzazione amministrativa</h2>"
                    "<p class='missing'>Unita amministrative non determinate: "
                    f"{html.escape(admin.note or 'nessuna sorgente disponibile')}.</p>")
        head = ["Comune", "ISTAT", "Codice catastale", "Provincia", "Regione",
                "Quota dell'area"]
        rows = [[unit.name, unit.istat_code, unit.cadastral_code, unit.province_code,
                 unit.region_name, f"{unit.area_share_pct:.1f} %"]
                for unit in admin.municipalities]
        note = ""
        if admin.is_multi_municipality:
            note = ("<p class='small'>L'area ricade in piu' Comuni: le elaborazioni "
                    "catastali sono suddivise per Comune.</p>")
        return "<h2>2. Localizzazione amministrativa</h2>" + self._table(head, rows) + note

    def _section_alerts(self) -> str:
        if not self.options.include_alerts:
            return ""
        alerts = self.report.sorted_alerts()
        if not alerts:
            return ("<h2>3. Segnalazioni</h2>"
                    "<p class='missing'>Nessuna segnalazione generata dalle regole attive.</p>")
        blocks = ["<h2>3. Segnalazioni</h2>"]
        for alert in alerts:
            reference = (f"<div class='small'>Riferimento indicato dalla fonte: "
                         f"{html.escape(alert.legal_reference)}</div>"
                         if alert.legal_reference else "")
            blocks.append(
                f"<div class='alert alert-{alert.level.value}'>"
                f"<div class='level'>{alert.level.label_it}</div>"
                f"<div><b>{html.escape(alert.title)}</b></div>"
                f"<div>{html.escape(alert.detail)}</div>"
                f"<div class='evidence'>Natura del dato: {alert.evidence_level.label_it}</div>"
                f"{reference}</div>")
        return "\n".join(blocks)

    def _section_cadastre(self) -> str:
        if not self.options.include_cadastre:
            return ""
        rows = self.report.cadastre
        if not rows:
            return ("<h2>4. Catasto</h2><p class='missing'>Nessuna particella catastale "
                    "disponibile per quest'area (servizio non disponibile oppure territorio "
                    "gestito dai catasti provinciali di Trento e Bolzano).</p>")
        blocks = ["<h2>4. Catasto</h2>"]
        blocks.append("<p class='small'>Le superfici indicate sono <b>grafiche</b>, "
                      "calcolate dalla geometria della mappa catastale: non corrispondono "
                      "necessariamente alle superfici censite in banca dati.</p>")
        grouped: Dict[str, List[CadastralRow]] = {}
        for row in rows:
            grouped.setdefault(row.municipality or "-", []).append(row)
        head = ["Foglio", "Particella", "Superficie grafica", "Superficie interessata",
                "% particella"]
        for municipality, items in grouped.items():
            total = sum(item.area_intersect_m2 for item in items)
            blocks.append(f"<h3>{html.escape(municipality)} "
                          f"({len(items)} particelle, {measure.format_area(total)})</h3>")
            body = [[item.sheet, item.parcel,
                     measure.format_area(item.area_cadastral_m2),
                     measure.format_area(item.area_intersect_m2),
                     f"{item.intersect_pct:.1f} %"]
                    for item in items[:self.options.max_cadastral_rows]]
            blocks.append(self._table(head, body))
            if len(items) > self.options.max_cadastral_rows:
                blocks.append(f"<p class='small'>Elenco troncato a "
                              f"{self.options.max_cadastral_rows} righe: la tabella completa "
                              f"e' nel file CSV/XLSX allegato.</p>")
        return "\n".join(blocks)

    def _knowledge_sections(self) -> List[str]:
        """One section per taxonomy category that belongs to the dossier."""
        blocks: List[str] = []
        index = 5
        # Sub-categories (a national park, an art. 136 designation) are reported inside
        # their dossier macro-section, not as extra sections.
        by_category: Dict[str, List[SourceResult]] = {}
        for raw_category, results in self.report.by_category().items():
            by_category.setdefault(self.taxonomy.root_of(raw_category), []).extend(results)
        for category in self.taxonomy.knowledge():
            if category.id == "cadastral":
                continue
            results = by_category.get(category.id, [])
            title = f"<h2>{index}. {html.escape(category.label('it'))}</h2>"
            index += 1
            if not results:
                blocks.append(title + "<p class='missing'>Nessuna sorgente configurata o "
                                      "disponibile per questa categoria nell'area.</p>")
                continue
            blocks.append(title + self._results_block(results))
        return blocks

    def _results_block(self, results: Sequence[SourceResult]) -> str:
        """Render the findings of one category."""
        parts: List[str] = []
        present = [r for r in results if r.present]
        empty = [r for r in results if r.ok and not r.present]
        failed = [r for r in results if not r.ok]
        for result in present:
            provenance = result.provenance
            lines = [f"<h3>{html.escape(result.source_name)}</h3>"]
            info = [
                f"Elementi interessati: {result.feature_count}",
                f"Superficie interessata: {measure.format_area(result.intersect_area_m2)} "
                f"({result.intersect_pct:.1f}% dell'area)"
                if result.intersect_area_m2 else "",
                f"Distanza minima: {measure.format_distance(result.min_distance_m)}"
                if result.min_distance_m else "",
                f"Ente: {provenance.authority}" if provenance and provenance.authority else "",
                f"Natura del dato: {result.evidence_level.label_it}",
                f"Aggiornamento dichiarato: {provenance.update_frequency}"
                if provenance and provenance.update_frequency else "",
                f"Scala/accuratezza: {provenance.scale or ''} "
                f"{('+/- ' + str(provenance.accuracy_m) + ' m') if provenance and provenance.accuracy_m else ''}"
                if provenance and (provenance.scale or provenance.accuracy_m) else "",
            ]
            lines.append("<ul>" + "".join(f"<li>{html.escape(item)}</li>"
                                          for item in info if item) + "</ul>")
            hits = [hit for hit in result.hits if hit.distance_m == 0.0][
                :self.options.max_features_per_source]
            if hits:
                head = ["Elemento", "Superficie interessata", "% area"]
                body = [[hit.label or hit.fid,
                         measure.format_area(hit.intersect_area_m2),
                         f"{hit.intersect_pct:.1f} %"] for hit in hits]
                lines.append(self._table(head, body))
            parts.append("\n".join(lines))
        if empty:
            names = ", ".join(html.escape(r.source_name) for r in empty)
            parts.append(f"<p class='small'>Fonti interrogate senza elementi nell'area: "
                         f"{names}.</p>")
        if failed:
            names = "; ".join(f"{html.escape(r.source_name)} ({html.escape(r.error[:90])})"
                              for r in failed)
            parts.append(f"<p class='missing'>Fonti non disponibili al momento "
                         f"dell'analisi: {names}. Il quadro di questa categoria e' "
                         f"incompleto.</p>")
        return "\n".join(parts)

    def _section_cultural_heritage(self) -> str:
        """Render the heritage section, when the module ran for this area.

        The section is built by the cultural-heritage module itself: this engine keeps
        owning the document, that module keeps owning what heritage data may be said to
        mean. The table renderer is handed over so the styling stays the document's.
        """
        if not self.options.include_cultural_heritage:
            return ""
        from .cultural_heritage.engine import outcome_of
        from .cultural_heritage.report import CulturalHeritageReport

        payload = self.report.module("cultural_heritage")
        if not payload:
            return ""
        outcome = outcome_of(self.report)
        if not outcome.themes:
            return ""
        number = 5 + len([c for c in self.taxonomy.knowledge() if c.id != "cadastral"])
        section = CulturalHeritageReport(outcome, number=number,
                                         max_rows=self.options.max_heritage_rows)
        return section.to_html(self._table)

    def _section_terrain(self) -> str:
        if not self.options.include_terrain:
            return ""
        terrain = self.report.terrain
        title = "<h2>Dati altimetrici</h2>"
        if terrain is None or not terrain.available:
            return title + "<p class='missing'>Dati altimetrici non disponibili.</p>"
        rows = [
            ("Quota minima", f"{terrain.elevation_min:.0f} m"),
            ("Quota massima", f"{terrain.elevation_max:.0f} m"),
            ("Quota media", f"{terrain.elevation_mean:.0f} m"),
            ("Quota mediana", f"{terrain.elevation_median:.0f} m"),
            ("Dislivello", f"{terrain.elevation_range:.0f} m"),
        ]
        if terrain.slope_mean_pct is not None:
            rows.append(("Pendenza media", f"{terrain.slope_mean_pct:.1f} %"))
            rows.append(("Pendenza massima", f"{terrain.slope_max_pct:.1f} %"))
        rows.append(("Risoluzione DEM", f"{terrain.cell_size_m:g} m"))
        blocks = [title, self._table_two_columns(rows)]
        if terrain.slope_classes:
            head = ["Classe di pendenza", "Superficie", "% area"]
            body = [[item.label, measure.format_area(item.area_m2), f"{item.area_pct:.1f} %"]
                    for item in terrain.slope_classes]
            blocks.append("<h3>Classi di pendenza</h3>" + self._table(head, body))
        if terrain.aspect_histogram:
            head = ["Esposizione", "% area"]
            body = [[name, f"{value:.1f} %"] for name, value in terrain.aspect_histogram.items()]
            blocks.append("<h3>Esposizione</h3>" + self._table(head, body))
        if terrain.provenance is not None:
            blocks.append(f"<p class='small'>Fonte DEM: "
                          f"{html.escape(terrain.provenance.source_name or terrain.source_id)} "
                          f"- {html.escape(terrain.provenance.attribution)}</p>")
        if terrain.note:
            blocks.append(f"<p class='small'>{html.escape(terrain.note)}</p>")
        return "\n".join(blocks)

    def _section_sources(self) -> str:
        if not self.options.include_sources:
            return ""
        head = ["Fonte", "Ente", "Stato", "Natura del dato", "Licenza", "Acquisito il"]
        rows = []
        for result in self.report.all_results:
            provenance = result.provenance
            rows.append([
                result.source_name,
                provenance.authority if provenance else "",
                result.status.value,
                result.evidence_level.label_it,
                provenance.license if provenance else "",
                self._format_date(provenance.retrieved_at if provenance else ""),
            ])
        if not rows:
            return "<h2>Fonti</h2><p class='missing'>Nessuna fonte interrogata.</p>"
        attributions = sorted({r.provenance.attribution for r in self.report.all_results
                               if r.provenance and r.provenance.attribution})
        extra = ""
        if attributions:
            extra = ("<p class='small'>Attribuzioni: " +
                     html.escape(" | ".join(attributions)) + "</p>")
        return "<h2>Fonti</h2>" + self._table(head, rows) + extra

    def _section_metadata(self) -> str:
        rows = [
            ("Plugin", f"{PLUGIN_NAME} {self.report.plugin_version or PLUGIN_VERSION}"),
            ("QGIS", self.report.qgis_version),
            ("Inizio analisi", self._format_date(self.report.started_at)),
            ("Fine analisi", self._format_date(self.report.finished_at)),
            ("Fonti interrogate", str(len(self.report.all_results))),
            ("Fonti non disponibili", str(len(self.report.sources_failed))),
            ("Layer prodotti", str(len(self.report.layers))),
        ]
        warnings = ""
        if self.report.warnings:
            items = "".join(f"<li>{html.escape(item)}</li>" for item in self.report.warnings)
            warnings = f"<h3>Avvisi di elaborazione</h3><ul>{items}</ul>"
        return "<h2>Metadati dell'elaborazione</h2>" + self._table_two_columns(rows) + warnings

    @staticmethod
    def _metric_bbox(area: Dict[str, Any]) -> Dict[str, float]:
        """Return the bounding box in the metric work CRS.

        The stored bbox is in the area CRS, which is often geographic: printing degrees
        with one decimal ("11.2 - 11.2") tells the reader nothing.
        """
        stored = area.get("bbox", {}) or {}
        fallback = {"min_x": 0.0, "min_y": 0.0, "max_x": 0.0, "max_y": 0.0,
                    "width": 0.0, "height": 0.0}
        wkt = area.get("geometry_wkt", "")
        work_crs = area.get("work_crs", "")
        if not wkt or not work_crs:
            return {**fallback, **stored}
        try:
            from qgis.core import QgsGeometry

            from ..core import crs as crs_utils
            from ..core import geometry as geom_utils

            geometry = QgsGeometry.fromWkt(wkt)
            rect = crs_utils.bbox_in(geometry, area.get("crs", crs_utils.WGS84), work_crs)
            return geom_utils.bbox_dict(rect)
        except Exception:  # pragma: no cover - defensive
            return {**fallback, **stored}

    # ------------------------------------------------------------------ tables

    @staticmethod
    def _table(head: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
        header = "".join(f"<th>{html.escape(str(item))}</th>" for item in head)
        body = "".join(
            "<tr>" + "".join(f"<td>{html.escape(str(cell))}</td>" for cell in row) + "</tr>"
            for row in rows)
        return f"<table><tr>{header}</tr>{body}</table>"

    @staticmethod
    def _table_two_columns(rows: Sequence[Sequence[Any]]) -> str:
        body = "".join(
            f"<tr><th style='width:35%'>{html.escape(str(key))}</th>"
            f"<td>{html.escape(str(value))}</td></tr>" for key, value in rows)
        return f"<table>{body}</table>"

    @staticmethod
    def _format_date(value: str) -> str:
        if not value:
            return ""
        try:
            return datetime.fromisoformat(value).strftime("%d/%m/%Y %H:%M")
        except ValueError:
            return value

    # ------------------------------------------------------------------ tabular exports

    def cadastral_rows(self) -> List[List[Any]]:
        """Return the cadastral table as plain rows (header first)."""
        head = ["Comune", "Codice ISTAT", "Codice catastale", "Foglio", "Particella",
                "Identificativo nazionale", "Superficie grafica (m2)",
                "Superficie interessata (m2)", "% particella interessata"]
        rows = [[row.municipality, row.istat_code, row.cadastral_code, row.sheet, row.parcel,
                 row.national_ref, round(row.area_cadastral_m2, 2),
                 round(row.area_intersect_m2, 2), round(row.intersect_pct, 2)]
                for row in self.report.cadastre]
        return [head] + rows

    def sources_rows(self) -> List[List[Any]]:
        """Return the source table as plain rows (header first)."""
        head = ["Fonte", "Id", "Ente", "Categoria", "Stato", "Elementi", "Natura del dato",
                "Licenza", "URL", "Acquisito il", "Errore"]
        rows = []
        for result in self.report.all_results:
            provenance = result.provenance
            rows.append([result.source_name, result.source_id,
                         provenance.authority if provenance else "", result.category,
                         result.status.value, result.feature_count,
                         result.evidence_level.label_it,
                         provenance.license if provenance else "",
                         provenance.url if provenance else "",
                         provenance.retrieved_at if provenance else "", result.error])
        return [head] + rows

    def alerts_rows(self) -> List[List[Any]]:
        """Return the alert table as plain rows (header first)."""
        head = ["Livello", "Titolo", "Dettaglio", "Categoria", "Natura del dato",
                "Riferimento", "Fonti"]
        rows = [[alert.level.label_it, alert.title, alert.detail, alert.category,
                 alert.evidence_level.label_it, alert.legal_reference,
                 ", ".join(alert.source_ids)]
                for alert in self.report.sorted_alerts()]
        return [head] + rows

    def write_csv(self, path: Path, rows: Sequence[Sequence[Any]]) -> Path:
        """Write a table as UTF-8 CSV with the semicolon separator used in Italy."""
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerows(rows)
        return path

    def _csv_fallback(self, path: Path, reason: str) -> None:
        """Write the tables as CSV when the XLSX writer cannot be used safely."""
        log.warning(f"{reason}: esporto le tabelle in CSV")
        base = path.with_suffix("")
        self.write_csv(Path(f"{base}_catasto.csv"), self.cadastral_rows())
        self.write_csv(Path(f"{base}_fonti.csv"), self.sources_rows())
        self.write_csv(Path(f"{base}_segnalazioni.csv"), self.alerts_rows())

    def write_xlsx(self, path: Path) -> Optional[Path]:
        """Write every table to one XLSX workbook (falls back to CSV when unavailable).

        Safety note (verified on QGIS 3.40 / Windows): openpyxl uses **lxml** when it is
        installed, and lxml carries its own libxml2, which collides with the libxml2 GDAL
        has already used to parse GML. The result is an access violation that kills QGIS in
        unrelated code. Setting ``OPENPYXL_LXML=False`` before importing openpyxl avoids
        it; if openpyxl was already imported elsewhere with lxml enabled, the tables are
        exported as CSV rather than risking a crash.
        """
        import os

        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        os.environ.setdefault("OPENPYXL_LXML", "False")
        try:
            import openpyxl
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:  # pragma: no cover - optional dependency
            self._csv_fallback(path, "openpyxl non disponibile")
            return None
        if getattr(openpyxl, "LXML", False):  # pragma: no cover - environment dependent
            self._csv_fallback(
                path, "openpyxl e' gia' stato caricato con il backend lxml (incompatibile "
                      "con GDAL: rischio di crash)")
            return None
        workbook = Workbook()
        sheets = (("Catasto", self.cadastral_rows()), ("Fonti", self.sources_rows()),
                  ("Segnalazioni", self.alerts_rows()))
        for index, (name, rows) in enumerate(sheets):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = name
            for row in rows:
                sheet.append(list(row))
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            widths: Dict[int, int] = {}
            for row in rows[:200]:
                for column, value in enumerate(row, start=1):
                    widths[column] = min(max(widths.get(column, 10), len(str(value)) + 2), 60)
            for column, width in widths.items():
                sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = width
        workbook.save(str(path))
        return path
